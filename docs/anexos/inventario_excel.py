# Inventario de cobertura: recorre cada hoja del Excel y describe que informacion
# contiene, para verificar que el modelo de datos la cubra por completo.
import openpyxl, os, sys

ruta = os.path.expanduser("~/Downloads/Alimentaci\u00f3n.xlsx")
wb = openpyxl.load_workbook(ruta, data_only=True)

for hoja in wb.sheetnames:
    ws = wb[hoja]
    filas = list(ws.iter_rows(values_only=True))
    usadas = [f for f in filas if any(c is not None and str(c).strip() != "" for c in f)]
    print("=" * 78)
    print(f"HOJA: {hoja}   filas con datos: {len(usadas)} de {ws.max_row}   columnas: {ws.max_column}")
    print("=" * 78)
    for i, f in enumerate(usadas[:34], 1):
        celdas = [("" if c is None else str(c))[:22] for c in f[:11]]
        while celdas and celdas[-1] == "":
            celdas.pop()
        if celdas:
            print(f"  {i:>3} | " + " | ".join(celdas))
    if len(usadas) > 34:
        print(f"  ... {len(usadas) - 34} filas mas")
    print()
